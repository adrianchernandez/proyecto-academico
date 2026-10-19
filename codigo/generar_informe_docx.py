"""Convierte un informe en Markdown (con figuras referenciadas) a Word (.docx).

El Markdown es la fuente de verdad y se versiona; el .docx se genera para leer y
compartir, y NO se versiona (contiene figuras con cifras de la operación).

Soporta lo que usan los informes del proyecto: títulos (#, ##, ###, ####),
párrafos, negrita/cursiva/código en línea, listas con viñetas y numeradas,
citas (>), tablas con tuberías, bloques de código y figuras ``![leyenda](ruta)``
(ruta relativa a la raíz del proyecto).

Uso::

    python codigo/generar_informe_docx.py informes/Informe_avance_01.md
    python codigo/generar_informe_docx.py            # todos los informes/*.md
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
from comun import BASE  # noqa: E402

INFORMES = BASE / "informes"
ANCHO_FIGURA = Cm(16)
RE_IMG = re.compile(r"^!\[(?P<cap>[^\]]*)\]\((?P<ruta>[^)]+)\)\s*$")
RE_INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)")


def _estilos(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.5)
        s.top_margin = s.bottom_margin = Cm(2.5)
    for nombre, tam in [("Title", 20), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 11.5)]:
        st = doc.styles[nombre]
        st.font.name = "Calibri"
        st.font.size = Pt(tam)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0x15, 0x19, 0x1B)


def _runs(par, texto: str) -> None:
    """Escribe `texto` en el párrafo interpretando **negrita**, *cursiva* y `código`."""
    for trozo in RE_INLINE.split(texto):
        if not trozo:
            continue
        if trozo.startswith("**") and trozo.endswith("**"):
            par.add_run(trozo[2:-2]).bold = True
        elif trozo.startswith("`") and trozo.endswith("`"):
            r = par.add_run(trozo[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        elif trozo.startswith("*") and trozo.endswith("*") and len(trozo) > 2:
            par.add_run(trozo[1:-1]).italic = True
        else:
            par.add_run(trozo)


def _sombrear(celda, color_hex: str) -> None:
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def _tabla(doc: Document, filas: list[list[str]]) -> None:
    ncol = max(len(f) for f in filas)
    t = doc.add_table(rows=len(filas), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, fila in enumerate(filas):
        for j in range(ncol):
            celda = t.cell(i, j)
            celda.text = ""
            p = celda.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            _runs(p, fila[j] if j < len(fila) else "")
            for r in p.runs:
                r.font.size = Pt(9)
                if i == 0:
                    r.bold = True
            if i == 0:
                _sombrear(celda, "E8EEF1")
    doc.add_paragraph()


def _figura(doc: Document, ruta: str, leyenda: str) -> None:
    archivo = (BASE / ruta).resolve()
    if not archivo.exists():
        p = doc.add_paragraph()
        p.add_run(f"[Figura no encontrada: {ruta}]").italic = True
        return
    doc.add_picture(str(archivo), width=ANCHO_FIGURA)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(leyenda)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x3C, 0x47, 0x4B)


def convertir(md: Path, destino: Path | None = None) -> Path:
    destino = destino or md.with_suffix(".docx")
    doc = Document()
    _estilos(doc)
    lineas = md.read_text(encoding="utf-8").splitlines()

    parrafo: list[str] = []
    tabla: list[list[str]] = []
    codigo: list[str] | None = None
    ultimo_lista = None    # último párrafo de lista, para continuaciones sangradas

    def cerrar_parrafo():
        nonlocal parrafo
        if parrafo:
            p = doc.add_paragraph()
            _runs(p, " ".join(s.strip() for s in parrafo))
            parrafo = []

    def cerrar_tabla():
        nonlocal tabla
        if tabla:
            _tabla(doc, tabla)
            tabla = []

    for linea in lineas:
        # --- bloques de código ------------------------------------------
        if linea.strip().startswith("```"):
            cerrar_parrafo(); cerrar_tabla()
            if codigo is None:
                codigo = []
            else:
                p = doc.add_paragraph()
                r = p.add_run("\n".join(codigo))
                r.font.name = "Consolas"
                r.font.size = Pt(9)
                codigo = None
            continue
        if codigo is not None:
            codigo.append(linea)
            continue

        s = linea.rstrip()
        # --- tablas ------------------------------------------------------
        if s.startswith("|"):
            cerrar_parrafo()
            celdas = [c.strip() for c in s.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in celdas if c):
                continue          # fila separadora
            tabla.append(celdas)
            continue
        cerrar_tabla()

        # --- vacías y reglas --------------------------------------------
        if not s.strip():
            cerrar_parrafo(); ultimo_lista = None
            continue
        if s.strip() == "---":
            cerrar_parrafo()
            continue

        # --- títulos ------------------------------------------------------
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            cerrar_parrafo()
            nivel = len(m.group(1))
            texto = m.group(2).strip()
            if nivel == 1:
                doc.add_heading(texto, 0)
            else:
                doc.add_heading(texto, nivel - 1)
            continue

        # --- figuras ------------------------------------------------------
        m = RE_IMG.match(s.strip())
        if m:
            cerrar_parrafo()
            _figura(doc, m.group("ruta"), m.group("cap"))
            continue

        # --- citas --------------------------------------------------------
        if s.startswith(">"):
            cerrar_parrafo()
            texto = s.lstrip("> ").strip()
            if not texto:
                continue
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            _runs(p, texto)
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x3C, 0x47, 0x4B)
            continue

        # --- listas -------------------------------------------------------
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", s)
        if m and (m.group(1) == "" or ultimo_lista is not None):
            cerrar_parrafo()
            estilo = "List Number" if m.group(2)[0].isdigit() else "List Bullet"
            if m.group(1):
                estilo += " 2"
            p = doc.add_paragraph(style=estilo)
            _runs(p, m.group(3))
            ultimo_lista = p
            continue
        if ultimo_lista is not None and s.startswith("  "):
            ultimo_lista.add_run(" ")
            _runs(ultimo_lista, s.strip())
            continue

        # --- párrafo normal ----------------------------------------------
        ultimo_lista = None
        parrafo.append(s)

    cerrar_parrafo(); cerrar_tabla()
    doc.save(destino)
    return destino


if __name__ == "__main__":
    objetivos = [Path(a) for a in sys.argv[1:]] or sorted(INFORMES.glob("*.md"))
    for md in objetivos:
        md = md if md.is_absolute() else BASE / md
        salida = convertir(md)
        print(f"{md.name} → {salida.name} ({salida.stat().st_size / 1024:.0f} KB)")
